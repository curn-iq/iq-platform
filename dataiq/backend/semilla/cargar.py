"""Carga inicial (seed) de DataIQ desde semilla/tecnicas.xlsx.

Uso, desde dataiq/backend:

    uv run python -m semilla.cargar            # carga o recarga las técnicas
    uv run python -m semilla.cargar --forzar   # reemplaza versiones hechas por personas

Cada técnica del archivo queda con una sola versión, la 1, publicada. Se puede
correr varias veces: borra las versiones anteriores de esas técnicas y las vuelve
a cargar. Si alguna versión la creó una persona (no la carga), se detiene para no
borrar trabajo real, salvo con --forzar.
"""

import asyncio
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.adaptadores.salida.persistencia.base_datos import SesionLocal
from app.adaptadores.salida.persistencia.modelos.catalogo import (
    CategoriaInstrumental,
    DispositivoMedico,
    EquipoBiomedico,
    Instrumental,
    InstrumentalAlias,
    Sutura,
)
from app.adaptadores.salida.persistencia.modelos.tecnicas import (
    Especialidad,
    ItemTecnica,
    ItemTecnicaComponente,
    PosicionMesa,
    Tecnica,
    TecnicaVersion,
    TecnicaVersionDispositivoMedico,
    TecnicaVersionEquipoBiomedico,
    TecnicaVersionSutura,
    Zona,
)
from app.adaptadores.salida.persistencia.modelos.usuarios import Usuario
from app.dominio.tecnicas import EstadoVersion
from app.dominio.usuarios import RolUsuario

ARCHIVO = Path(__file__).with_name("tecnicas.xlsx")
CORREO_CARGA = "carga@dataiq.local"
SIN_DATO = {None, "", "—"}


# ---------- Lectura del archivo ----------


@dataclass
class Objeto:
    mesa: str | None
    numero: int | None
    texto: str
    componentes: list[str]
    celdas: list[tuple[int, int]]


@dataclass
class TecnicaArchivo:
    nombre: str
    especialidad: str
    datos: dict
    suturas: list[tuple[str, str, str | None, str | None]] = field(default_factory=list)
    equipos: list[str] = field(default_factory=list)
    dispositivos: list[str] = field(default_factory=list)
    tamanos: dict[str, tuple[int, int]] = field(default_factory=dict)
    objetos: list[Objeto] = field(default_factory=list)


def _filas(libro, hoja):
    hoja = libro[hoja]
    columnas = [c.value for c in hoja[1]]
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if any(v not in SIN_DATO for v in fila):
            yield dict(zip(columnas, fila, strict=False))


def _texto(valor):
    return None if valor in SIN_DATO else str(valor).strip()


def leer_archivo(ruta: Path):
    libro = load_workbook(ruta, read_only=True, data_only=True)
    tecnicas: dict[str, TecnicaArchivo] = {}
    for f in _filas(libro, "tecnicas"):
        tecnicas[f["nombre"]] = TecnicaArchivo(
            nombre=f["nombre"],
            especialidad=f["especialidad"],
            datos={
                "codigo_cups": _texto(f["codigo_cups"]),
                "descripcion": _texto(f["para_que_sirve"]),
                "indicaciones": _texto(f["indicaciones"]),
                "complicaciones": _texto(f["complicaciones"]),
                "tecnica_quirurgica": _texto(f["tecnica_quirurgica"]),
                "anestesia": _texto(f["anestesia"]),
                "posicion_paciente": _texto(f["posicion_paciente"]),
                "ropa": _texto(f["ropa"]),
                "fuente": _texto(f["fuente"]),
            },
        )
    for f in _filas(libro, "suturas"):
        tecnicas[f["tecnica"]].suturas.append(
            (f["plano"], f["sutura"], _texto(f["calibre"]), _texto(f["aguja"]))
        )
    for f in _filas(libro, "equipos"):
        tecnicas[f["tecnica"]].equipos.append(f["equipo"])
    for f in _filas(libro, "dispositivos"):
        tecnicas[f["tecnica"]].dispositivos.append(f["dispositivo"])
    for f in _filas(libro, "mesas"):
        if f["filas"] not in SIN_DATO:
            tecnicas[f["tecnica"]].tamanos[f["mesa"]] = (
                int(f["filas"]),
                int(f["columnas"]),
            )
    for f in _filas(libro, "objetos"):
        celdas = []
        for par in (_texto(f["celdas"]) or "").split(";"):
            if par.strip():
                fila, columna = par.split(",")
                celdas.append((int(fila), int(columna)))
        componentes = [c.strip() for c in (_texto(f["componentes"]) or "").split("+")]
        tecnicas[f["tecnica"]].objetos.append(
            Objeto(
                mesa=_texto(f["mesa"]),
                numero=int(f["numero"]) if f["numero"] not in SIN_DATO else None,
                texto=str(f["texto"]),
                componentes=[c for c in componentes if c],
                celdas=celdas,
            )
        )
    nuevos = [
        (f["nombre"], f["categoria"], _texto(f["alias"]))
        for f in _filas(libro, "catalogo_nuevo")
    ]
    return tecnicas, nuevos


def validar(tecnicas, instrumental_conocido: set[str]) -> list[str]:
    """Lo que la base rechazaría o que no tendría sentido cargar."""
    errores = []
    for t in tecnicas.values():
        ocupadas, numeros = {}, set()
        for o in t.objetos:
            for c in o.componentes:
                if c not in instrumental_conocido:
                    errores.append(f"{t.nombre}: «{c}» no está en el catálogo")
            if o.mesa is None:
                if o.numero is not None or o.celdas:
                    errores.append(
                        f"{t.nombre}: «{o.texto}» tiene número o celda sin mesa"
                    )
                continue
            if o.mesa not in t.tamanos:
                errores.append(
                    f"{t.nombre}: «{o.texto}» está en {o.mesa}, que no se lleva"
                )
                continue
            if (o.mesa, o.numero) in numeros:
                errores.append(f"{t.nombre}: número {o.numero} repetido en {o.mesa}")
            numeros.add((o.mesa, o.numero))
            filas, columnas = t.tamanos[o.mesa]
            for celda in o.celdas:
                if not (1 <= celda[0] <= filas and 1 <= celda[1] <= columnas):
                    errores.append(
                        f"{t.nombre}: {o.mesa} n.° {o.numero} fuera de la mesa"
                    )
                if (o.mesa, celda) in ocupadas:
                    errores.append(
                        f"{t.nombre}: {o.mesa} celda {celda} ocupada dos veces"
                    )
                ocupadas[(o.mesa, celda)] = o.numero
    return errores


# ---------- Escritura en la base ----------


async def _obtener_o_crear(sesion: AsyncSession, modelo, **campos):
    consulta = select(modelo)
    for nombre, valor in campos.items():
        columna = getattr(modelo, nombre)
        # En Sutura, dos campos vacíos cuentan como iguales (NULLS NOT DISTINCT)
        consulta = consulta.where(columna.is_not_distinct_from(valor))
    fila = (await sesion.scalars(consulta)).first()
    if fila is None:
        fila = modelo(**campos)
        sesion.add(fila)
        await sesion.flush()
    return fila


async def _usuario_de_carga(sesion: AsyncSession) -> Usuario:
    usuario = (
        await sesion.scalars(select(Usuario).where(Usuario.email == CORREO_CARGA))
    ).first()
    if usuario is None:
        usuario = Usuario(
            nombre="Carga inicial",
            email=CORREO_CARGA,
            password_hash="!",  # no corresponde a ninguna contraseña: no puede ingresar
            fecha_aceptacion_politica=datetime.now(UTC),
            version_politica="no aplica",
        )
        sesion.add(usuario)
    # Publica lo que carga, así que necesita un rol que publique
    usuario.rol = RolUsuario.revisor
    await sesion.flush()
    return usuario


async def _borrar_versiones(sesion: AsyncSession, versiones: list) -> None:
    ids = [v.id for v in versiones]
    if not ids:
        return
    items = select(ItemTecnica.id).where(ItemTecnica.version_id.in_(ids))
    await sesion.execute(delete(PosicionMesa).where(PosicionMesa.version_id.in_(ids)))
    await sesion.execute(
        delete(ItemTecnicaComponente).where(ItemTecnicaComponente.item_id.in_(items))
    )
    await sesion.execute(delete(ItemTecnica).where(ItemTecnica.version_id.in_(ids)))
    for tabla in (
        TecnicaVersionSutura,
        TecnicaVersionEquipoBiomedico,
        TecnicaVersionDispositivoMedico,
    ):
        await sesion.execute(delete(tabla).where(tabla.version_id.in_(ids)))
    await sesion.execute(delete(TecnicaVersion).where(TecnicaVersion.id.in_(ids)))


async def _borrar_huerfanos(sesion: AsyncSession) -> None:
    """Suturas, equipos y dispositivos que ninguna técnica usa ya."""
    usadas = select(TecnicaVersionSutura.sutura_id).union(
        select(ItemTecnicaComponente.sutura_id).where(
            ItemTecnicaComponente.sutura_id.is_not(None)
        )
    )
    await sesion.execute(delete(Sutura).where(Sutura.id.not_in(usadas)))
    await sesion.execute(
        delete(EquipoBiomedico).where(
            EquipoBiomedico.id.not_in(select(TecnicaVersionEquipoBiomedico.equipo_id))
        )
    )
    await sesion.execute(
        delete(DispositivoMedico).where(
            DispositivoMedico.id.not_in(
                select(TecnicaVersionDispositivoMedico.dispositivo_id)
            )
        )
    )


async def cargar(forzar: bool = False) -> None:
    tecnicas, nuevos = leer_archivo(ARCHIVO)

    async with SesionLocal() as sesion, sesion.begin():
        # Catálogo: los instrumentos nuevos del archivo, con su categoría y sus alias
        for nombre, categoria, alias in nuevos:
            cat = await _obtener_o_crear(
                sesion, CategoriaInstrumental, nombre=categoria
            )
            inst = (
                await sesion.scalars(
                    select(Instrumental).where(Instrumental.nombre == nombre)
                )
            ).first()
            if inst is None:
                inst = Instrumental(nombre=nombre, categoria_id=cat.id)
                sesion.add(inst)
                await sesion.flush()
            for a in (alias or "").split(";"):
                if a.strip():
                    await _obtener_o_crear(
                        sesion,
                        InstrumentalAlias,
                        instrumental_id=inst.id,
                        alias=a.strip(),
                    )
        instrumental = {
            i.nombre: i.id for i in (await sesion.scalars(select(Instrumental))).all()
        }

        errores = validar(tecnicas, set(instrumental))
        if errores:
            raise SystemExit("No se cargó nada:\n  " + "\n  ".join(errores))

        carga = await _usuario_de_carga(sesion)
        zonas = {
            nombre: (await _obtener_o_crear(sesion, Zona, nombre=nombre)).id
            for nombre in ("Mesa de Mayo", "Mesa de reserva")
        }
        ahora = datetime.now(UTC)

        for t in tecnicas.values():
            esp = await _obtener_o_crear(sesion, Especialidad, nombre=t.especialidad)
            tecnica = (
                await sesion.scalars(
                    select(Tecnica).where(
                        Tecnica.nombre == t.nombre, Tecnica.especialidad_id == esp.id
                    )
                )
            ).first()
            if tecnica is None:
                tecnica = Tecnica(
                    nombre=t.nombre, especialidad_id=esp.id, creado_por=carga.id
                )
                sesion.add(tecnica)
                await sesion.flush()

            versiones = (
                await sesion.scalars(
                    select(TecnicaVersion).where(
                        TecnicaVersion.tecnica_id == tecnica.id
                    )
                )
            ).all()
            de_personas = [v for v in versiones if v.creado_por != carga.id]
            cargas_viejas = (
                await sesion.scalars(
                    select(Usuario.id).where(Usuario.email.like("%@dataiq.local"))
                )
            ).all()
            de_personas = [v for v in de_personas if v.creado_por not in cargas_viejas]
            if de_personas and not forzar:
                raise SystemExit(
                    f"{t.nombre} tiene versiones creadas por personas; no se borran. "
                    "Usar --forzar para reemplazarlas."
                )
            await _borrar_versiones(sesion, versiones)

            version = TecnicaVersion(
                tecnica_id=tecnica.id,
                numero=1,
                estado=EstadoVersion.publicada,
                creado_por=carga.id,
                revisado_por=carga.id,
                fecha_revision=ahora,
                **t.datos,
            )
            sesion.add(version)
            await sesion.flush()

            # Una misma sutura puede usarse en varios planos: un solo registro con todos
            planos = defaultdict(list)
            for plano, nombre, calibre, aguja in t.suturas:
                planos[(nombre, calibre, aguja)].append(plano)
            for (nombre, calibre, aguja), usos in planos.items():
                sutura = await _obtener_o_crear(
                    sesion, Sutura, nombre=nombre, calibre=calibre, tipo_aguja=aguja
                )
                sesion.add(
                    TecnicaVersionSutura(
                        version_id=version.id, sutura_id=sutura.id, uso="; ".join(usos)
                    )
                )
            for nombre in dict.fromkeys(t.equipos):
                equipo = await _obtener_o_crear(sesion, EquipoBiomedico, nombre=nombre)
                sesion.add(
                    TecnicaVersionEquipoBiomedico(
                        version_id=version.id, equipo_id=equipo.id
                    )
                )
            for nombre in dict.fromkeys(t.dispositivos):
                disp = await _obtener_o_crear(sesion, DispositivoMedico, nombre=nombre)
                sesion.add(
                    TecnicaVersionDispositivoMedico(
                        version_id=version.id, dispositivo_id=disp.id
                    )
                )

            for o in t.objetos:
                item = ItemTecnica(
                    version_id=version.id,
                    zona_id=zonas[o.mesa] if o.mesa else None,
                    numero_leyenda=o.numero,
                    texto_fuente=o.texto,
                )
                sesion.add(item)
                await sesion.flush()
                for c in o.componentes:
                    sesion.add(
                        ItemTecnicaComponente(
                            item_id=item.id, instrumental_id=instrumental[c]
                        )
                    )
                for fila, columna in o.celdas:
                    sesion.add(
                        PosicionMesa(
                            item_id=item.id,
                            version_id=version.id,
                            zona_id=zonas[o.mesa],
                            fila=fila,
                            columna=columna,
                        )
                    )
            await sesion.flush()

        await _borrar_huerfanos(sesion)

    print(f"Cargadas {len(tecnicas)} técnicas desde {ARCHIVO.name}.")


if __name__ == "__main__":
    asyncio.run(cargar(forzar="--forzar" in sys.argv))
