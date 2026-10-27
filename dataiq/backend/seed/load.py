"""Carga inicial (seed) de DataIQ desde seed/tecnicas.xlsx.

Uso, desde dataiq/backend:

    uv run python -m seed.load            # carga o recarga las técnicas
    uv run python -m seed.load --force    # reemplaza versiones hechas por personas

Cada técnica del archivo queda con una sola versión, la 1, publicada. Se puede
correr varias veces: borra las versiones anteriores de esas técnicas y las vuelve
a cargar. Si alguna versión la creó una persona (no la carga), se detiene para no
borrar trabajo real, salvo con --force.
"""

import asyncio
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.adapters.outbound.persistence.database import SessionLocal
from app.adapters.outbound.persistence.models.catalog import (
    BiomedicalEquipment,
    Instrument,
    InstrumentAlias,
    InstrumentCategory,
    MedicalDevice,
    Suture,
)
from app.adapters.outbound.persistence.models.techniques import (
    ItemCell,
    ItemComponent,
    Specialty,
    Technique,
    TechniqueItem,
    TechniqueVersion,
    VersionDevice,
    VersionEquipment,
    VersionRejection,
    VersionSuture,
    VersionZone,
    Zone,
)
from app.adapters.outbound.persistence.models.users import User
from app.domain.techniques import VersionState
from app.domain.users import Role

SOURCE_FILE = Path(__file__).with_name("tecnicas.xlsx")
SEED_EMAIL = "carga@dataiq.local"
NO_DATA = {None, "", "—"}


# ---------- Lectura del archivo ----------


@dataclass
class SeedItem:
    stand: str | None
    number: int | None
    source_text: str
    components: list[str]
    cells: list[tuple[int, int]]


@dataclass
class SeedTechnique:
    name: str
    specialty: str
    data: dict
    sutures: list[tuple[str, str, str | None, str | None]] = field(default_factory=list)
    equipment: list[str] = field(default_factory=list)
    devices: list[str] = field(default_factory=list)
    sizes: dict[str, tuple[int, int]] = field(default_factory=dict)
    # Mesas que la técnica no lleva (columna no_lleva = «Sí»)
    unused: list[str] = field(default_factory=list)
    items: list[SeedItem] = field(default_factory=list)


def _rows(workbook, sheet):
    sheet = workbook[sheet]
    columns = [c.value for c in sheet[1]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(v not in NO_DATA for v in row):
            yield dict(zip(columns, row, strict=False))


def _text(value_):
    return None if value_ in NO_DATA else str(value_).strip()


def read_file(path_: Path):
    workbook = load_workbook(path_, read_only=True, data_only=True)
    techniques: dict[str, SeedTechnique] = {}
    for f in _rows(workbook, "tecnicas"):
        techniques[f["nombre"]] = SeedTechnique(
            name=f["nombre"],
            specialty=f["especialidad"],
            data={
                "cups_code": _text(f["codigo_cups"]),
                "description": _text(f["para_que_sirve"]),
                "indications": _text(f["indicaciones"]),
                "complications": _text(f["complicaciones"]),
                "surgical_technique": _text(f["tecnica_quirurgica"]),
                "anesthesia": _text(f["anestesia"]),
                "patient_position": _text(f["posicion_paciente"]),
                "drapes": _text(f["ropa"]),
                "source": _text(f["fuente"]),
            },
        )
    for f in _rows(workbook, "suturas"):
        techniques[f["tecnica"]].sutures.append(
            (f["plano"], f["sutura"], _text(f["calibre"]), _text(f["aguja"]))
        )
    for f in _rows(workbook, "equipos"):
        techniques[f["tecnica"]].equipment.append(f["equipo"])
    for f in _rows(workbook, "dispositivos"):
        techniques[f["tecnica"]].devices.append(f["dispositivo"])
    for f in _rows(workbook, "mesas"):
        if f["no_lleva"] == "Sí":
            techniques[f["tecnica"]].unused.append(f["mesa"])
        elif f["filas"] not in NO_DATA:
            techniques[f["tecnica"]].sizes[f["mesa"]] = (
                int(f["filas"]),
                int(f["columnas"]),
            )
    for f in _rows(workbook, "objetos"):
        cells = []
        for pair in (_text(f["celdas"]) or "").split(";"):
            if pair.strip():
                row, column = pair.split(",")
                cells.append((int(row), int(column)))
        components = [c.strip() for c in (_text(f["componentes"]) or "").split("+")]
        techniques[f["tecnica"]].items.append(
            SeedItem(
                stand=_text(f["mesa"]),
                number=int(f["numero"]) if f["numero"] not in NO_DATA else None,
                source_text=str(f["texto"]),
                components=[c for c in components if c],
                cells=cells,
            )
        )
    # El catálogo base (hoja catalogo) y lo que se le agregó (catalogo_nuevo):
    # así la carga también sirve en una base de datos vacía.
    instruments = [
        (f["nombre"], f["categoria"], _text(f["alias"]), _text(f["descripcion"]))
        for f in _rows(workbook, "catalogo")
    ] + [
        (f["nombre"], f["categoria"], _text(f["alias"]), None)
        for f in _rows(workbook, "catalogo_nuevo")
    ]
    return techniques, instruments


def validate(techniques, known_instruments: set[str]) -> list[str]:
    """Lo que la base rechazaría o que no tendría sentido cargar."""
    errors = []
    for t in techniques.values():
        occupied, numbers = {}, set()
        for o in t.items:
            for c in o.components:
                if c not in known_instruments:
                    errors.append(f"{t.name}: «{c}» no está en el catálogo")
            if o.stand is None:
                if o.number is not None or o.cells:
                    errors.append(
                        f"{t.name}: «{o.source_text}» tiene número o celda sin mesa"
                    )
                continue
            if o.stand not in t.sizes:
                errors.append(
                    f"{t.name}: «{o.source_text}» está en {o.stand}, que no se lleva"
                )
                continue
            if (o.stand, o.number) in numbers:
                errors.append(f"{t.name}: número {o.number} repetido en {o.stand}")
            numbers.add((o.stand, o.number))
            rows, columns = t.sizes[o.stand]
            for cell in o.cells:
                if not (1 <= cell[0] <= rows and 1 <= cell[1] <= columns):
                    errors.append(
                        f"{t.name}: {o.stand} n.° {o.number} fuera de la mesa"
                    )
                if (o.stand, cell) in occupied:
                    errors.append(f"{t.name}: {o.stand} celda {cell} ocupada dos veces")
                occupied[(o.stand, cell)] = o.number
    return errors


# ---------- Escritura en la base ----------


async def _get_or_create(session: AsyncSession, model, **fields):
    query = select(model)
    for name, value_ in fields.items():
        column = getattr(model, name)
        # En sutures, dos campos vacíos cuentan como iguales (NULLS NOT DISTINCT)
        query = query.where(column.is_not_distinct_from(value_))
    row = (await session.scalars(query)).first()
    if row is None:
        row = model(**fields)
        session.add(row)
        await session.flush()
    return row


async def _seed_user(session: AsyncSession) -> User:
    user = (await session.scalars(select(User).where(User.email == SEED_EMAIL))).first()
    if user is None:
        user = User(
            name="Carga inicial",
            email=SEED_EMAIL,
            password_hash="!",  # no corresponde a ninguna contraseña: no puede ingresar
            policy_accepted_at=datetime.now(UTC),
            policy_version="no aplica",
        )
        session.add(user)
    # Publica lo que carga, así que necesita un rol que publique
    user.role = Role.reviewer
    await session.flush()
    return user


async def _delete_versions(session: AsyncSession, versions: list) -> None:
    ids = [v.id for v in versions]
    if not ids:
        return
    items = select(TechniqueItem.id).where(TechniqueItem.version_id.in_(ids))
    await session.execute(delete(ItemCell).where(ItemCell.version_id.in_(ids)))
    await session.execute(delete(ItemComponent).where(ItemComponent.item_id.in_(items)))
    await session.execute(
        delete(TechniqueItem).where(TechniqueItem.version_id.in_(ids))
    )
    for table in (
        VersionRejection,
        VersionZone,
        VersionSuture,
        VersionEquipment,
        VersionDevice,
    ):
        await session.execute(delete(table).where(table.version_id.in_(ids)))
    await session.execute(delete(TechniqueVersion).where(TechniqueVersion.id.in_(ids)))


async def _delete_orphans(session: AsyncSession) -> None:
    """Suturas, equipos y dispositivos que ninguna técnica usa ya."""
    used = select(VersionSuture.suture_id).union(
        select(ItemComponent.suture_id).where(ItemComponent.suture_id.is_not(None))
    )
    await session.execute(delete(Suture).where(Suture.id.not_in(used)))
    await session.execute(
        delete(BiomedicalEquipment).where(
            BiomedicalEquipment.id.not_in(select(VersionEquipment.equipment_id))
        )
    )
    await session.execute(
        delete(MedicalDevice).where(
            MedicalDevice.id.not_in(select(VersionDevice.device_id))
        )
    )


async def load(force: bool = False) -> None:
    techniques, catalog = read_file(SOURCE_FILE)

    async with SessionLocal() as session, session.begin():
        # Catálogo: los instrumentos que falten, con su categoría y sus alias. De
        # los que ya existen solo se llena la descripción si no tienen (un revisor
        # pudo haberlos editado).
        for name, category, alias, description in catalog:
            cat = await _get_or_create(session, InstrumentCategory, name=category)
            aliases = [a.strip() for a in (alias or "").split(";") if a.strip()]
            inst = (
                await session.scalars(select(Instrument).where(Instrument.name == name))
            ).first()
            if inst is None and aliases:
                # Un nombre corregido: el instrumento conserva su id y su nombre
                # anterior queda como alias.
                inst = (
                    await session.scalars(
                        select(Instrument).where(Instrument.name.in_(aliases))
                    )
                ).first()
                if inst is not None:
                    inst.name = name
            if inst is None:
                inst = Instrument(name=name, category_id=cat.id)
                session.add(inst)
                await session.flush()
            if description and not inst.description:
                inst.description = description
            for a in aliases:
                await _get_or_create(
                    session, InstrumentAlias, instrument_id=inst.id, alias=a
                )
        instruments = {
            i.name: i.id for i in (await session.scalars(select(Instrument))).all()
        }

        errors = validate(techniques, set(instruments))
        if errors:
            raise SystemExit("No se cargó nada:\n  " + "\n  ".join(errors))

        seed_user = await _seed_user(session)
        zones = {
            name: (await _get_or_create(session, Zone, name=name)).id
            for name in ("Mesa de Mayo", "Mesa de reserva")
        }
        now_ = datetime.now(UTC)

        for t in techniques.values():
            spec = await _get_or_create(session, Specialty, name=t.specialty)
            technique = (
                await session.scalars(
                    select(Technique).where(
                        Technique.name == t.name, Technique.specialty_id == spec.id
                    )
                )
            ).first()
            if technique is None:
                technique = Technique(
                    name=t.name, specialty_id=spec.id, created_by=seed_user.id
                )
                session.add(technique)
                await session.flush()

            versions = (
                await session.scalars(
                    select(TechniqueVersion).where(
                        TechniqueVersion.technique_id == technique.id
                    )
                )
            ).all()
            by_people = [v for v in versions if v.created_by != seed_user.id]
            old_loads = (
                await session.scalars(
                    select(User.id).where(User.email.like("%@dataiq.local"))
                )
            ).all()
            by_people = [v for v in by_people if v.created_by not in old_loads]
            if by_people and not force:
                raise SystemExit(
                    f"{t.name} tiene versiones creadas por personas; no se borran. "
                    "Usar --force para reemplazarlas."
                )
            await _delete_versions(session, versions)

            version = TechniqueVersion(
                technique_id=technique.id,
                number=1,
                state=VersionState.published,
                created_by=seed_user.id,
                reviewed_by=seed_user.id,
                reviewed_at=now_,
                **t.data,
            )
            session.add(version)
            await session.flush()

            for stand, (rows, columns) in t.sizes.items():
                session.add(
                    VersionZone(
                        version_id=version.id,
                        zone_id=zones[stand],
                        used=True,
                        grid_rows=rows,
                        grid_columns=columns,
                    )
                )
            for stand in t.unused:
                session.add(
                    VersionZone(version_id=version.id, zone_id=zones[stand], used=False)
                )

            # Una misma sutura puede usarse en varios planos: un solo registro con todos
            planes = defaultdict(list)
            for plane, name, gauge, needle in t.sutures:
                planes[(name, gauge, needle)].append(plane)
            for (name, gauge, needle), usages in planes.items():
                suture = await _get_or_create(
                    session, Suture, name=name, gauge=gauge, needle_type=needle
                )
                session.add(
                    VersionSuture(
                        version_id=version.id,
                        suture_id=suture.id,
                        usage="; ".join(usages),
                    )
                )
            for name in dict.fromkeys(t.equipment):
                equipment_item = await _get_or_create(
                    session, BiomedicalEquipment, name=name
                )
                session.add(
                    VersionEquipment(
                        version_id=version.id, equipment_id=equipment_item.id
                    )
                )
            for name in dict.fromkeys(t.devices):
                disp = await _get_or_create(session, MedicalDevice, name=name)
                session.add(VersionDevice(version_id=version.id, device_id=disp.id))

            for o in t.items:
                item = TechniqueItem(
                    version_id=version.id,
                    zone_id=zones[o.stand] if o.stand else None,
                    legend_number=o.number,
                    source_text=o.source_text,
                )
                session.add(item)
                await session.flush()
                for c in o.components:
                    session.add(
                        ItemComponent(item_id=item.id, instrument_id=instruments[c])
                    )
                for row, column in o.cells:
                    session.add(
                        ItemCell(
                            item_id=item.id,
                            version_id=version.id,
                            zone_id=zones[o.stand],
                            grid_row=row,
                            grid_column=column,
                        )
                    )
            await session.flush()

        await _delete_orphans(session)

    print(f"Cargadas {len(techniques)} técnicas desde {SOURCE_FILE.name}.")


if __name__ == "__main__":
    asyncio.run(load(force="--force" in sys.argv))
